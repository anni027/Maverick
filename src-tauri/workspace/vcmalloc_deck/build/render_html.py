"""Render deck.html slides to high-resolution PNGs, then assemble a PPTX.

Pipeline:  HTML/CSS  ->  headless Chrome (Playwright)  ->  PNG  ->  PPTX
Slides are rasterised at 2x for crisp output on 13.333x7.5in slides.
"""
import os
import sys
import glob
import shutil

from playwright.sync_api import sync_playwright

CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

SLIDE_W, SLIDE_H = 1280, 720
SCALE = 2                      # 2560x1440 output
OUTDIR = 'render'


def render_slides(html='deck.html', outdir=OUTDIR, scale=SCALE, only=None):
    exe = CHROME if os.path.exists(CHROME) else EDGE
    if os.path.isdir(outdir):
        for f in glob.glob(os.path.join(outdir, '*.png')):
            os.remove(f)
    os.makedirs(outdir, exist_ok=True)

    url = 'file:///' + os.path.abspath(html).replace('\\', '/')
    made = []
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=exe, args=[
            '--force-color-profile=srgb', '--disable-lcd-text',
            '--font-render-hinting=none'])
        pg = b.new_page(viewport={'width': SLIDE_W + 40,
                                  'height': SLIDE_H + 40},
                        device_scale_factor=scale)
        pg.goto(url, wait_until='networkidle')
        pg.wait_for_timeout(700)
        # ensure every image (charts) has decoded
        pg.wait_for_function(
            "Array.from(document.images).every(i => i.complete)", timeout=30000)
        pg.add_style_tag(content=(
            "body{background:#fff;margin:0;padding:0}"
            ".slide{margin:0 !important;box-shadow:none !important}"
            "*{ -webkit-print-color-adjust:exact; print-color-adjust:exact;}"
        ))
        pg.wait_for_timeout(300)

        slides = pg.query_selector_all('.slide')
        print('slides found:', len(slides))
        for i, el in enumerate(slides, 1):
            if only and i not in only:
                continue
            out = os.path.join(outdir, 'slide%02d.png' % i)
            el.screenshot(path=out, scale='device')
            made.append(out)
            print('  rendered', out)
        b.close()
    return made


def build_pptx(pngs, out='VCMalloc.pptx', title='VCMalloc'):
    from pptx import Presentation
    from pptx.util import Emu, Inches
    prs = Presentation()
    prs.slide_width = Emu(12192000)     # 13.333in
    prs.slide_height = Emu(6858000)     # 7.5in
    blank = prs.slide_layouts[6]
    for p in pngs:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(p, 0, 0, width=prs.slide_width,
                             height=prs.slide_height)
    prs.core_properties.title = title
    prs.core_properties.author = 'VCMalloc deck generator'
    prs.core_properties.subject = ('VCMalloc: A Virtually Contiguous Memory '
                                   'Allocator')
    prs.save(out)
    print('saved', out, len(pngs), 'slides',
          round(os.path.getsize(out) / 1e6, 2), 'MB')
    return out


if __name__ == '__main__':
    only = None
    if len(sys.argv) > 2:
        only = set(int(x) for x in sys.argv[2].split(','))
    html = sys.argv[1] if len(sys.argv) > 1 else 'deck.html'
    imgs = render_slides(html, only=only)
    if not only:
        build_pptx(sorted(glob.glob(os.path.join(OUTDIR, '*.png'))))
